"""
Application layer orchestrator for the Graph/SharePoint flow.
Mirrors process_weekly_metrics.py (local), but receives already-read
data instead of file paths — since BytesIO from Graph can't be re-read
like a local Path can.
"""
import io
import openpyxl
from domain.model import WeekData, MasterEmployee
from domain.errors import ATError
from adapters.excel_write import create_workbook, create_empty_week_block
from infrastructure.config import OUTPUT_SHEET_NAME

from domain.business_rules import (
    match_employees,
    apply_business_rules,
    detect_cross_month,
    build_placeholder_employees, 
    split_cross_month_range
)
from adapters.sharepoint_client import (
    GraphContext,
    get_file_by_path,
    upload_file_graph,
    build_output_path,
    build_next_month_output_path,
)

def handle_cross_month_remote(context: GraphContext, weekly_data: WeekData, master_employees: list[MasterEmployee],) -> bool:
    placeholders = build_placeholder_employees(master_employees)

    closing_start, closing_end, opening_start, opening_end = split_cross_month_range(
        weekly_data.week_start, weekly_data.week_end
    )
    closing_folder, closing_file_name = build_output_path(closing_start.year, closing_start.month)

    opening_folder, opening_file_name = build_next_month_output_path(
        weekly_data.week_start.year, weekly_data.week_start.month
    )

    blocks = [
        (closing_folder, closing_file_name, closing_start, closing_end),
        (opening_folder, opening_file_name, opening_start, opening_end),
    ]

    any_created = False

    for folder, file_name, block_start, block_end in blocks:
        path = f"{folder}/{file_name}"
        existing_bytes = get_file_by_path(context, path)

        if existing_bytes is not None:
            workbook = openpyxl.load_workbook(existing_bytes)
            worksheet = workbook[OUTPUT_SHEET_NAME]
        else:
            workbook, worksheet = create_workbook()

        created = create_empty_week_block(worksheet, block_start, block_end, placeholders)
        if created:
            any_created = True

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        upload_file_graph(context, folder, file_name, buffer.getvalue())

    return any_created


    
def process_metric_remote(
    context: GraphContext,
    weekly_data: WeekData,
    master_employees: list[MasterEmployee],
) -> WeekData:

    cross_month = detect_cross_month(weekly_data.week_start, weekly_data.week_end)

    if cross_month:
        templates_created = handle_cross_month_remote(context, weekly_data, master_employees)

        if templates_created:
            raise ATError(
                "ERR013",
                "The weekly file crosses two months — templates created for manual completion"
            )
        else:
            raise ATError(
                "ERR015",
                "The weekly file crosses two months — templates already existed, nothing changed"
            )

    matched_employees = match_employees(
        weekly_data.employees,
        master_employees
    )

    weekly_data.employees = apply_business_rules(
        matched_employees
    )

    return weekly_data